import os
import io
import re
import streamlit as st
from pymongo import MongoClient
from bson import ObjectId
from datetime import datetime, timedelta, timezone
from dotenv import load_dotenv
import requests
import json
from google_auth_oauthlib.flow import Flow
import bcrypt
from urllib.parse import quote_plus, urlencode
import certifi
import logging
import secrets
from prometheus_client import Counter, Histogram, start_http_server, REGISTRY

# Document processing imports (graceful fallback if not installed)
try:
    import pypdf
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    import docx
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    import openpyxl
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False

try:
    from duckduckgo_search import DDGS
    WEB_SEARCH_SUPPORT = True
except ImportError:
    WEB_SEARCH_SUPPORT = False

# ----------------- Logging Setup -----------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger("chatbot")

def get_metric(name, metric_type, *args, **kwargs):
    if name in REGISTRY._names_to_collectors:
        return REGISTRY._names_to_collectors[name]
    else:
        return metric_type(name, *args, **kwargs)

# ----------------- Prometheus Metrics -----------------
REQUEST_COUNT = get_metric('chatbot_requests_total', Counter, 'Total requests')
LOGIN_COUNT = get_metric('chatbot_logins_total', Counter, 'Total logins')
REGISTER_COUNT = get_metric('chatbot_registers_total', Counter, 'Total registrations')
CHAT_COUNT = get_metric('chatbot_chats_total', Counter, 'Total chat messages')
ERROR_COUNT = get_metric('chatbot_errors_total', Counter, 'Total errors')
CHAT_LATENCY = get_metric('chatbot_chat_latency_seconds', Histogram, 'Chat response latency')

if not hasattr(st, "prometheus_metrics_server_started"):
    try:
        start_http_server(8000)
        logger.info("Prometheus metrics server started on port 8000")
    except Exception as e:
        logger.warning(f"Prometheus metrics server could not start : {e}")
    st.prometheus_metrics_server_started = True

# ----------------- Load environment -----------------
load_dotenv()
DEFAULT_API_KEY = os.getenv("KRUTRIM_API_KEY")
API_URL = "https://cloud.olakrutrim.com/v1/chat/completions"

# ----------------- Available Models -----------------
AVAILABLE_MODELS = [
    {"id": "Krutrim-spectre-v2", "name": "Krutrim Spectre v2", "description": "Krutrim's flagship model"},
    {"id": "Meta-Llama-3.1-8B-Instruct", "name": "Llama 3.1 8B", "description": "Fast & efficient"},
    {"id": "Meta-Llama-3.1-70B-Instruct", "name": "Llama 3.1 70B", "description": "Powerful reasoning"},
    {"id": "DeepSeek-R1", "name": "DeepSeek R1", "description": "Advanced reasoning"},
]

# ----------------- Google OAuth -----------------
SCOPES = [
    "openid",
    "https://www.googleapis.com/auth/userinfo.email",
    "https://www.googleapis.com/auth/userinfo.profile",
]

def get_google_oauth_url():
    """Generate Google OAuth authorization URL"""
    client_id = os.getenv("GOOGLE_CLIENT_ID")
    if not client_id:
        logger.error("Google Client ID not configured")
        return None
    state = secrets.token_urlsafe(32)
    st.session_state.oauth_state = state
    redirect_uri = os.getenv("OAUTH_REDIRECT_URI", "http://localhost:8501")
    params = {
        "client_id": client_id,
        "redirect_uri": redirect_uri,
        "scope": " ".join(SCOPES),
        "response_type": "code",
        "state": state,
        "access_type": "offline",
        "prompt": "select_account"
    }
    oauth_url = "https://accounts.google.com/o/oauth2/auth?" + urlencode(params)
    return oauth_url

def exchange_code_for_token(code, state):
    """Exchange authorization code for access token and get user info"""
    if state != st.session_state.get("oauth_state"):
        logger.error("OAuth state mismatch - possible CSRF attack")
        return None
    client_id = os.getenv("GOOGLE_CLIENT_ID")
    client_secret = os.getenv("GOOGLE_CLIENT_SECRET")
    redirect_uri = os.getenv("OAUTH_REDIRECT_URI", "http://localhost:8501")
    if not client_id or not client_secret:
        logger.error("Google OAuth credentials not configured")
        return None
    token_data = {
        "client_id": client_id,
        "client_secret": client_secret,
        "code": code,
        "grant_type": "authorization_code",
        "redirect_uri": redirect_uri
    }
    try:
        token_response = requests.post("https://oauth2.googleapis.com/token", data=token_data)
        if token_response.status_code != 200:
            logger.error(f"Token exchange failed: {token_response.status_code} {token_response.text}")
            return None
        token_info = token_response.json()
        access_token = token_info.get("access_token")
        if not access_token:
            logger.error("No access token received")
            return None
        user_response = requests.get(
            "https://www.googleapis.com/oauth2/v2/userinfo",
            headers={"Authorization": f"Bearer {access_token}"}
        )
        if user_response.status_code != 200:
            logger.error(f"Failed to fetch user info: {user_response.status_code} {user_response.text}")
            return None
        user_info = user_response.json()
        logger.info(f"Successfully retrieved user info for: {user_info.get('email')}")
        return user_info
    except Exception as e:
        logger.error(f"OAuth token exchange failed: {e}")
        return None

def create_or_update_google_user(user_info):
    """Create or update user in database with Google profile info"""
    google_id = user_info.get("id")
    email = user_info.get("email")
    name = user_info.get("name")
    picture = user_info.get("picture")
    if not email or not google_id:
        logger.error("Missing required user info from Google")
        return None
    user = users_col.find_one({"$or": [{"google_id": google_id}, {"email": email}]})
    if user:
        logger.info(f"Updating existing user: {email}")
        update_data = {
            "google_id": google_id,
            "name": name,
            "picture": picture,
            "last_login": datetime.now(timezone.utc)
        }
        users_col.update_one({"_id": user["_id"]}, {"$set": update_data})
        user = users_col.find_one({"_id": user["_id"]})
    else:
        logger.info(f"Creating new Google user: {email}")
        new_user = {
            "email": email,
            "google_id": google_id,
            "name": name,
            "picture": picture,
            "api_key": None,
            "tokens_used_today": 0,
            "last_reset": datetime.now(timezone.utc),
            "created_at": datetime.now(timezone.utc),
            "last_login": datetime.now(timezone.utc),
            "auth_method": "google"
        }
        result = users_col.insert_one(new_user)
        user = users_col.find_one({"_id": result.inserted_id})
        REGISTER_COUNT.inc()
    return user

# ----------------- Chat Management Functions -----------------
def create_new_chat(user_id):
    """Create a new chat session"""
    try:
        chat_metadata_col = db.chat_metadata
        chat_doc = {
            "user_id": user_id,
            "title": "New Chat",
            "created_at": datetime.now(timezone.utc),
            "updated_at": datetime.now(timezone.utc),
            "message_count": 0
        }
        result = chat_metadata_col.insert_one(chat_doc)
        return str(result.inserted_id)
    except Exception as e:
        logger.error(f"Error creating new chat: {e}")
        return None

def get_user_chats(user_id):
    """Get all chats for a user"""
    try:
        chat_metadata_col = db.chat_metadata
        chats = list(chat_metadata_col.find(
            {"user_id": user_id},
            {"title": 1, "created_at": 1, "updated_at": 1, "message_count": 1}
        ).sort("updated_at", -1))
        return chats
    except Exception as e:
        logger.error(f"Error getting user chats: {e}")
        return []

def get_chat_messages(chat_id):
    """Get messages for a specific chat"""
    try:
        messages = list(chats_col.find(
            {"chat_id": ObjectId(chat_id)},
            {"role": 1, "content": 1, "timestamp": 1}
        ).sort("timestamp", 1))
        return [(msg["role"], msg["content"]) for msg in messages]
    except Exception:
        return []

def save_chat_message(user_id, chat_id, role, content):
    """Save a message to a specific chat"""
    try:
        message_doc = {
            "user_id": user_id,
            "chat_id": ObjectId(chat_id),
            "role": role,
            "content": content,
            "timestamp": datetime.now(timezone.utc)
        }
        chats_col.insert_one(message_doc)
        chat_metadata_col = db.chat_metadata
        if role == "user" and len(content) > 0:
            title = content[:50] + "..." if len(content) > 50 else content
            chat_metadata_col.update_one(
                {"_id": ObjectId(chat_id)},
                {
                    "$set": {"title": title, "updated_at": datetime.now(timezone.utc)},
                    "$inc": {"message_count": 1}
                },
                upsert=True
            )
        else:
            chat_metadata_col.update_one(
                {"_id": ObjectId(chat_id)},
                {"$set": {"updated_at": datetime.now(timezone.utc)}},
                upsert=True
            )
    except Exception as e:
        logger.error(f"Error saving chat message: {e}")

def delete_chat(chat_id):
    """Delete a chat and all its messages"""
    try:
        chat_metadata_col = db.chat_metadata
        chats_col.delete_many({"chat_id": ObjectId(chat_id)})
        chat_metadata_col.delete_one({"_id": ObjectId(chat_id)})
        return True
    except Exception as e:
        logger.error(f"Error deleting chat: {e}")
        return False

# ----------------- Group Chat Functions -----------------
def create_group_chat(creator_id, name, member_emails):
    """Create a group chat with a list of member emails"""
    try:
        group_chats_col = db.group_chats
        members = [creator_id]
        for email in member_emails:
            email = email.strip()
            if email:
                u = users_col.find_one({"email": email}, {"_id": 1})
                if u:
                    members.append(u["_id"])
        group_doc = {
            "name": name,
            "creator_id": creator_id,
            "members": members,
            "member_emails": [email.strip() for email in member_emails if email.strip()],
            "created_at": datetime.now(timezone.utc),
            "updated_at": datetime.now(timezone.utc),
        }
        result = group_chats_col.insert_one(group_doc)
        return str(result.inserted_id)
    except Exception as e:
        logger.error(f"Error creating group chat: {e}")
        return None

def get_user_group_chats(user_id):
    """Get group chats where user is a member"""
    try:
        group_chats_col = db.group_chats
        groups = list(group_chats_col.find(
            {"members": user_id},
            {"name": 1, "member_emails": 1, "created_at": 1, "updated_at": 1}
        ).sort("updated_at", -1))
        return groups
    except Exception as e:
        logger.error(f"Error getting group chats: {e}")
        return []

def add_member_to_group(group_id, email):
    """Add a member to a group chat by email"""
    try:
        group_chats_col = db.group_chats
        u = users_col.find_one({"email": email.strip()}, {"_id": 1})
        update = {"$addToSet": {"member_emails": email.strip()}, "$set": {"updated_at": datetime.now(timezone.utc)}}
        if u:
            update["$addToSet"]["members"] = u["_id"]
        group_chats_col.update_one({"_id": ObjectId(group_id)}, update)
        return True
    except Exception as e:
        logger.error(f"Error adding member to group: {e}")
        return False

# ----------------- Document Processing Functions -----------------
def extract_text_from_pdf(file_bytes):
    """Extract text from PDF bytes"""
    if not PDF_SUPPORT:
        return "PDF support not available. Install pypdf."
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        texts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                texts.append(t)
        return "\n".join(texts)
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return f"Error extracting PDF: {e}"

def extract_text_from_docx(file_bytes):
    """Extract text from DOCX bytes"""
    if not DOCX_SUPPORT:
        return "DOCX support not available. Install python-docx."
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return f"Error extracting DOCX: {e}"

def extract_text_from_excel(file_bytes):
    """Extract text from Excel bytes"""
    if not EXCEL_SUPPORT:
        return "Excel support not available. Install openpyxl."
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                row_str = "\t".join([str(cell) if cell is not None else "" for cell in row])
                if row_str.strip():
                    parts.append(row_str)
        return "\n".join(parts)
    except Exception as e:
        logger.error(f"Excel extraction error: {e}")
        return f"Error extracting Excel: {e}"

def extract_text_from_file(uploaded_file):
    """Detect file type and extract text, returning a dict"""
    try:
        file_bytes = uploaded_file.read()
        name = uploaded_file.name
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if ext == "pdf":
            content = extract_text_from_pdf(file_bytes)
        elif ext == "docx":
            content = extract_text_from_docx(file_bytes)
        elif ext in ("xlsx", "xls"):
            content = extract_text_from_excel(file_bytes)
        else:
            try:
                content = file_bytes.decode("utf-8")
            except Exception:
                content = "Unable to decode file content."
        return {"filename": name, "content": content, "type": ext}
    except Exception as e:
        logger.error(f"File extraction error: {e}")
        return {"filename": uploaded_file.name, "content": f"Error: {e}", "type": "unknown"}

def get_document_context(docs, max_chars=3000):
    """Build context string from uploaded documents"""
    parts = []
    total = 0
    for doc in docs:
        header = f"--- Document: {doc['filename']} ---\n"
        content = doc.get("content", "")
        remaining = max_chars - total - len(header)
        if remaining <= 0:
            break
        chunk = content[:remaining]
        parts.append(header + chunk + "\n")
        total += len(header) + len(chunk)
    return "\n".join(parts)

# ----------------- Web Search Function -----------------
def web_search(query, max_results=5):
    """Search the web using DuckDuckGo and return formatted results"""
    if not WEB_SEARCH_SUPPORT:
        return ""
    try:
        with DDGS() as ddgs:
            results = list(ddgs.text(query, max_results=max_results))
        formatted = []
        for i, r in enumerate(results, 1):
            formatted.append(
                f"{i}. **{r.get('title', '')}**\n{r.get('body', '')}\nSource: {r.get('href', '')}"
            )
        return "\n\n".join(formatted)
    except Exception as e:
        logger.error(f"Web search error: {e}")
        return ""

# ----------------- Code Block Extraction -----------------
CANVAS_LANGUAGES = ["python", "javascript", "typescript", "java", "cpp", "c", "bash", "sql", "json", "html", "css", "text"]
def extract_code_blocks(text):
    """Extract all code blocks from markdown text"""
    pattern = r'```(\w*)\s*\n?(.*?)```'
    matches = re.findall(pattern, text, re.DOTALL)
    return [(lang or 'text', code.strip()) for lang, code in matches]

# ----------------- MongoDB -----------------
@st.cache_resource
def get_mongo_client():
    """Initializes and returns a cached MongoDB client."""
    logger.info("Connecting to MongoDB...")
    username = quote_plus("chatbot")
    mongo_pass = os.getenv("MONGO_PASS")
    if not mongo_pass:
        raise ValueError("MONGO_PASS environment variable is not set")
    password = quote_plus(mongo_pass)
    dbname = os.getenv("MONGO_DB")
    mongo_uri = f"mongodb+srv://{username}:{password}@cluster0.57nirib.mongodb.net/{dbname}?retryWrites=true&w=majority&tls=true"
    client = MongoClient(mongo_uri, tlsCAFile=certifi.where(), serverSelectionTimeoutMS=10000)
    try:
        client.admin.command("ping")
        logger.info("MongoDB connected!")
    except Exception as e:
        logger.error(f"MongoDB connection failed: {e}")
    return client

client = get_mongo_client()
db = client[os.getenv("MONGO_DB")]
users_col = db.users
chats_col = db.chats

# ----------------- Authentication -----------------
def hash_password(password):
    logger.info("Hashing password")
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt())

def check_password(password, hashed):
    logger.info("Checking password")
    return bcrypt.checkpw(password.encode(), hashed)

def register_user(email, password):
    logger.info(f"Registering user: {email}")
    REGISTER_COUNT.inc()
    if users_col.find_one({"email": email}):
        logger.warning(f"User already exists: {email}")
        return False, "User already exists!"
    hashed = hash_password(password)
    users_col.insert_one({
        "email": email,
        "password_hash": hashed,
        "api_key": None,
        "tokens_used_today": 0,
        "last_reset": datetime.now(timezone.utc),
        "created_at": datetime.now(timezone.utc),
        "auth_method": "email"
    })
    logger.info(f"User registered: {email}")
    return True, "Registered successfully!"

def login_user(email, password):
    logger.info(f"Login attempt: {email}")
    user = users_col.find_one({"email": email})
    if not user:
        logger.warning(f"User not found: {email}")
        ERROR_COUNT.inc()
        return False, "User not found."
    if not user.get("password_hash"):
        logger.warning(f"User {email} was created via Google OAuth, cannot login with password")
        ERROR_COUNT.inc()
        return False, "This account was created with Google. Please use 'Sign in with Google'."
    if check_password(password, user["password_hash"]):
        LOGIN_COUNT.inc()
        logger.info(f"Login successful: {email}")
        users_col.update_one({"_id": user["_id"]}, {"$set": {"last_login": datetime.now(timezone.utc)}})
        return True, user
    logger.warning(f"Incorrect password for user: {email}")
    ERROR_COUNT.inc()
    return False, "Incorrect password."

# ----------------- Rate Limiting -----------------
def can_use_tokens(user, tokens_needed):
    logger.info(f"Checking token usage for user: {user['email']}")
    last_reset = user.get("last_reset")
    if not isinstance(last_reset, datetime):
        last_reset = datetime.now(timezone.utc)
    else:
        if last_reset.tzinfo is None:
            last_reset = last_reset.replace(tzinfo=timezone.utc)
    if datetime.now(timezone.utc) - last_reset > timedelta(days=1):
        logger.info(f"Resetting tokens for user: {user['email']}")
        users_col.update_one(
            {"_id": user["_id"]},
            {"$set": {"tokens_used_today": 0, "last_reset": datetime.now(timezone.utc)}}
        )
        user["tokens_used_today"] = 0
    allowed = user["tokens_used_today"] + tokens_needed <= 2000
    logger.info(f"Token allowed: {allowed} for user: {user['email']}")
    return allowed

def increment_tokens(user, tokens):
    logger.info(f"Incrementing tokens by {tokens} for user: {user['email']}")
    users_col.update_one({"_id": user["_id"]}, {"$inc": {"tokens_used_today": tokens}})

# ----------------- Krutrim Chat -----------------
def chat_with_krutrim(messages, api_key=None, model="Krutrim-spectre-v2", max_tokens=2048):
    logger.info(f"Sending chat request to Krutrim API with model: {model}")
    key_to_use = api_key if api_key else DEFAULT_API_KEY
    headers = {"Authorization": f"Bearer {key_to_use}", "Content-Type": "application/json"}
    payload = {"model": model, "messages": messages, "max_tokens": max_tokens, "temperature": 0.7}
    try:
        with CHAT_LATENCY.time():
            resp = requests.post(API_URL, headers=headers, json=payload)
        REQUEST_COUNT.inc()
        if resp.status_code == 200:
            data = resp.json()
            logger.info("Received response from Krutrim API")
            return data["choices"][0]["message"]["content"], payload["max_tokens"]
        logger.error(f"Krutrim API error: {resp.status_code}, {resp.text}")
        ERROR_COUNT.inc()
        return f"❌ Error: {resp.status_code}, {resp.text}", 0
    except Exception as e:
        logger.error(f"Exception during chat API call: {e}")
        ERROR_COUNT.inc()
        return f"❌ Exception: {str(e)}", 0

# ----------------- Streamlit UI -----------------
st.set_page_config(
    page_title="Krutrim Chatbot",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ----------------- Session State -----------------
if "authenticated" not in st.session_state:
    st.session_state.authenticated = False
if "user_email" not in st.session_state:
    st.session_state.user_email = None
if "user" not in st.session_state:
    st.session_state.user = None
if "messages" not in st.session_state:
    st.session_state.messages = []
if "oauth_state" not in st.session_state:
    st.session_state.oauth_state = None
if "is_generating" not in st.session_state:
    st.session_state.is_generating = False
if "current_chat_id" not in st.session_state:
    st.session_state.current_chat_id = None
if "show_profile" not in st.session_state:
    st.session_state.show_profile = False
if "chat_history" not in st.session_state:
    st.session_state.chat_history = {}
# New feature session state
if "selected_model" not in st.session_state:
    st.session_state.selected_model = "Krutrim-spectre-v2"
if "web_search_enabled" not in st.session_state:
    st.session_state.web_search_enabled = False
if "uploaded_docs" not in st.session_state:
    st.session_state.uploaded_docs = []
if "canvas_content" not in st.session_state:
    st.session_state.canvas_content = None
if "show_canvas" not in st.session_state:
    st.session_state.show_canvas = False
if "canvas_language" not in st.session_state:
    st.session_state.canvas_language = "python"
if "chat_mode" not in st.session_state:
    st.session_state.chat_mode = "normal"
if "current_group_id" not in st.session_state:
    st.session_state.current_group_id = None
if "show_group_manager" not in st.session_state:
    st.session_state.show_group_manager = False

# ----------------- OAuth Callback -----------------
query_params = st.query_params
if "code" in query_params and "state" in query_params and not st.session_state.authenticated:
    logger.info("Processing OAuth callback")
    code = query_params.get("code")
    state = query_params.get("state")
    user_info = exchange_code_for_token(code, state)
    if user_info:
        user = create_or_update_google_user(user_info)
        if user:
            st.session_state.authenticated = True
            st.session_state.user_email = user.get("email")
            st.session_state.user = user
            st.session_state.messages = []
            LOGIN_COUNT.inc()
            logger.info(f"Google user successfully logged in: {user.get('email')}")
            st.query_params.clear()
            st.rerun()
        else:
            st.error("Failed to create or update user account")
    else:
        st.error("Failed to process Google sign-in")
        st.query_params.clear()

# ----------------- CSS -----------------
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');

    /* Global Styles */
    .stApp {
        background: linear-gradient(135deg, #0f0f23 0%, #1a1a2e 50%, #16213e 100%);
        font-family: 'Inter', sans-serif;
    }

    /* Login Container */
    .login-container {
        background: rgba(255, 255, 255, 0.05);
        backdrop-filter: blur(20px);
        border: 1px solid rgba(255, 255, 255, 0.1);
        border-radius: 20px;
        padding: 2.5rem;
        margin: 2rem auto;
        max-width: 450px;
        box-shadow: 0 20px 60px rgba(0, 0, 0, 0.3);
    }

    /* Navbar */
    .navbar {
        position: fixed;
        top: 0; left: 0; right: 0;
        background: rgba(15, 15, 35, 0.9);
        backdrop-filter: blur(20px);
        color: #fff;
        padding: 1rem 2rem;
        z-index: 1000;
        box-shadow: 0 4px 20px rgba(0,0,0,0.2);
        display: flex;
        justify-content: space-between;
        align-items: center;
        border-bottom: 1px solid rgba(255, 255, 255, 0.1);
    }
    .navbar-title {
        font-size: 1.4rem;
        font-weight: 700;
        letter-spacing: 1px;
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        background-clip: text;
    }

    /* Main app padding */
    .main .block-container {
        padding-top: 5rem;
    }

    /* Sidebar */
    section[data-testid="stSidebar"] {
        background: rgba(15, 15, 35, 0.9);
        backdrop-filter: blur(20px);
        border-right: 1px solid rgba(255, 255, 255, 0.1);
    }

    /* Buttons */
    .stButton > button {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: #fff;
        border: none;
        border-radius: 12px;
        padding: 0.7rem 1.5rem;
        font-size: 1rem;
        font-weight: 500;
        transition: all 0.3s ease;
        box-shadow: 0 4px 15px rgba(102, 126, 234, 0.2);
        width: 100%;
        margin: 0.25rem 0;
    }
    .stButton > button:hover {
        transform: translateY(-2px);
        box-shadow: 0 8px 25px rgba(102, 126, 234, 0.4);
        background: linear-gradient(135deg, #764ba2 0%, #667eea 100%);
    }
    .stButton > button:active {
        transform: translateY(0);
    }

    /* Form Elements */
    .stTextInput > div > div > input,
    .stTextArea > div > div > textarea,
    input[type="password"],
    input[type="email"] {
        background: rgba(255, 255, 255, 0.05);
        border: 1px solid rgba(255, 255, 255, 0.1);
        border-radius: 12px;
        color: #fff;
        padding: 0.8rem 1rem;
        font-size: 1rem;
        transition: all 0.3s ease;
    }
    .stTextInput > div > div > input:focus,
    input[type="password"]:focus,
    input[type="email"]:focus {
        border-color: #667eea;
        box-shadow: 0 0 0 2px rgba(102, 126, 234, 0.2);
        outline: none;
    }

    /* Chat Input */
    .stChatInput {
        position: sticky;
        bottom: 0;
        background: rgba(15, 15, 35, 0.95);
        backdrop-filter: blur(20px);
        padding: 20px 0;
        margin-top: 20px;
        border-top: 1px solid rgba(255, 255, 255, 0.1);
    }
    .stChatInput > div {
        max-width: 900px;
        margin: 0 auto;
    }
    .stChatInput > div > div {
        background: rgba(64, 65, 79, 0.3);
        border-radius: 26px;
        border: 1px solid rgba(255, 255, 255, 0.1);
        box-shadow: 0 4px 20px rgba(0, 0, 0, 0.1);
        transition: all 0.3s ease;
    }
    .stChatInput > div > div:focus-within {
        border-color: rgba(102, 126, 234, 0.5);
        box-shadow: 0 4px 30px rgba(102, 126, 234, 0.2);
        background: rgba(64, 65, 79, 0.5);
    }
    .stChatInput > div > div > input {
        background: transparent !important;
        border: none !important;
        color: #ffffff !important;
        padding: 16px 24px !important;
        font-size: 16px !important;
        line-height: 1.5 !important;
        min-height: 24px !important;
    }
    .stChatInput > div > div > input::placeholder {
        color: rgba(255, 255, 255, 0.5) !important;
    }

    /* Success/Error Messages */
    .stSuccess { background: rgba(52, 168, 83, 0.1); border: 1px solid rgba(52, 168, 83, 0.3); border-radius: 12px; }
    .stError { background: rgba(234, 67, 53, 0.1); border: 1px solid rgba(234, 67, 53, 0.3); border-radius: 12px; }
    .stInfo { background: rgba(66, 133, 244, 0.1); border: 1px solid rgba(66, 133, 244, 0.3); border-radius: 12px; }

    /* Chat Messages */
    .user-message {
        background: #2f2f2f;
        color: #ffffff;
        border-radius: 18px;
        padding: 12px 16px;
        margin: 8px 0 16px auto;
        max-width: 80%;
        word-wrap: break-word;
        box-shadow: 0 1px 2px rgba(0, 0, 0, 0.1);
        border: 1px solid rgba(255, 255, 255, 0.1);
    }
    .assistant-message {
        background: rgba(247, 247, 248, 0.05);
        color: #ffffff;
        border-radius: 18px;
        padding: 16px 20px;
        margin: 8px auto 16px 0;
        max-width: 85%;
        word-wrap: break-word;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.1);
        border: 1px solid rgba(255, 255, 255, 0.08);
        line-height: 1.6;
    }
    .message-container { display: flex; align-items: flex-start; margin: 12px 0; gap: 12px; }
    .user-container { justify-content: flex-end; }
    .assistant-container { justify-content: flex-start; }
    .message-avatar {
        width: 32px; height: 32px; border-radius: 50%;
        display: flex; align-items: center; justify-content: center;
        font-size: 14px; flex-shrink: 0; margin-top: 4px;
    }
    .user-avatar { background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; order: 2; }
    .assistant-avatar { background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%); color: white; }

    /* Typing Indicator */
    .typing-indicator {
        display: flex; align-items: center; gap: 8px;
        color: rgba(255, 255, 255, 0.6); font-style: italic;
        padding: 12px 16px; background: rgba(255, 255, 255, 0.03);
        border-radius: 18px; margin: 8px auto 16px 0; max-width: fit-content;
    }
    .typing-dots { display: flex; gap: 3px; }
    .typing-dot {
        width: 6px; height: 6px; border-radius: 50%;
        background: rgba(255, 255, 255, 0.4);
        animation: typing 1.4s infinite;
    }
    .typing-dot:nth-child(2) { animation-delay: 0.2s; }
    .typing-dot:nth-child(3) { animation-delay: 0.4s; }
    @keyframes typing {
        0%, 60%, 100% { opacity: 0.3; transform: scale(0.8); }
        30% { opacity: 1; transform: scale(1); }
    }

    /* Typography */
    h1, h2, h3 { color: #fff; font-weight: 600; }
    h1 {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        -webkit-background-clip: text; -webkit-text-fill-color: transparent;
        background-clip: text; text-align: center; margin-bottom: 2rem;
    }

    /* Animations */
    @keyframes fadeIn { from { opacity: 0; transform: translateY(20px); } to { opacity: 1; transform: translateY(0); } }
    .fade-in { animation: fadeIn 0.6s ease-out; }
    .loading-dots::after { content: '.'; animation: dots 1.5s steps(5, end) infinite; }
    @keyframes dots {
        0%, 20% { color: rgba(255,255,255,0); text-shadow: .25em 0 0 rgba(255,255,255,0), .5em 0 0 rgba(255,255,255,0); }
        40% { color: white; text-shadow: .25em 0 0 rgba(255,255,255,0), .5em 0 0 rgba(255,255,255,0); }
        60% { text-shadow: .25em 0 0 white, .5em 0 0 rgba(255,255,255,0); }
        80%, 100% { text-shadow: .25em 0 0 white, .5em 0 0 white; }
    }

    /* Scrollbar */
    ::-webkit-scrollbar { width: 8px; }
    ::-webkit-scrollbar-track { background: rgba(255, 255, 255, 0.05); }
    ::-webkit-scrollbar-thumb { background: rgba(102, 126, 234, 0.5); border-radius: 4px; }
    ::-webkit-scrollbar-thumb:hover { background: rgba(102, 126, 234, 0.8); }

    /* Stats Cards */
    .stats-card {
        background: rgba(255, 255, 255, 0.05); border-radius: 12px; padding: 1rem;
        margin: 0.5rem 0; border: 1px solid rgba(255, 255, 255, 0.1); text-align: center;
    }
    .stats-number { font-size: 1.5rem; font-weight: 700; color: #667eea; }
    .stats-label { font-size: 0.9rem; color: rgba(255, 255, 255, 0.7); margin-top: 0.25rem; }

    /* Model Badge */
    .model-badge {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        border-radius: 8px; padding: 2px 8px; font-size: 0.75rem; color: white; margin-left: 8px;
    }

    /* Canvas Panel */
    .canvas-panel {
        background: rgba(30, 30, 50, 0.9);
        border-left: 1px solid rgba(255,255,255,0.1);
        border-radius: 12px;
        padding: 1rem;
    }

    /* Feature Toolbar */
    .feature-toolbar {
        display: flex; gap: 8px; align-items: center;
        padding: 8px 0; border-top: 1px solid rgba(255,255,255,0.1);
        margin-bottom: 8px;
    }

    /* Doc Chip */
    .doc-chip {
        background: rgba(102,126,234,0.2);
        border: 1px solid rgba(102,126,234,0.4);
        border-radius: 16px; padding: 4px 10px; font-size: 0.8rem;
        color: #a0aec0; display: inline-flex; align-items: center; gap: 4px; margin: 2px;
    }

    /* Fixed header */
    .fixed-header {
        position: fixed; top: 0; left: 0; right: 0; z-index: 1000;
        background: rgba(15, 15, 35, 0.95); backdrop-filter: blur(20px);
        border-bottom: 1px solid rgba(255,255,255,0.1); padding: 0.8rem 2rem;
    }
    .header-content {
        display: flex; justify-content: space-between; align-items: center;
        max-width: 1200px; margin: 0 auto;
    }
    .header-title { font-size: 1.4rem; font-weight: 700; color: #667eea; }
    .header-spacer { height: 70px; }
    </style>
""", unsafe_allow_html=True)

# ----------------- Navigation Bar -----------------
if st.session_state.authenticated:
    st.markdown("""
        <style>
        .stButton > button {
            width: 100%; border-radius: 8px; border: 1px solid rgba(255, 255, 255, 0.2);
            background: rgba(255, 255, 255, 0.1); color: white; transition: all 0.2s ease;
        }
        .stButton > button:hover { background: rgba(255, 255, 255, 0.2); border-color: rgba(102, 126, 234, 0.5); }
        </style>
        <div class="fixed-header">
            <div class="header-content">
                <div class="header-title">🤖 Krutrim AI</div>
            </div>
        </div>
        <div class="header-spacer"></div>
    """, unsafe_allow_html=True)

    col1, col2, col3 = st.columns([8, 1, 1])
    with col2:
        if st.button("👤 Profile", help="View profile settings", key="nav_profile"):
            st.session_state.show_profile = not st.session_state.show_profile
            st.rerun()
    with col3:
        if st.button("🚪 Logout", help="Sign out", key="nav_logout"):
            logger.info(f"User logged out: {st.session_state.user['email']}")
            st.session_state.authenticated = False
            st.session_state.user_email = None
            st.session_state.user = None
            st.session_state.messages = []
            st.session_state.current_chat_id = None
            st.session_state.show_profile = False
            st.rerun()

# =============================================================================
# LOGIN / REGISTRATION
# =============================================================================
if not st.session_state.authenticated:
    logger.info("Rendering login/register UI")
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        st.markdown('<div class="login-container fade-in">', unsafe_allow_html=True)
        st.markdown("""
            <h1 style="text-align: center; margin-bottom: 2rem; font-size: 2.5rem;">
                🤖 Welcome to YAHANAR AI
            </h1>
            <p style="text-align: center; color: rgba(255,255,255,0.7); font-size: 1.1rem; margin-bottom: 2rem;">
                Your intelligent assistant powered by advanced AI
            </p>
        """, unsafe_allow_html=True)

        choice = st.radio("Choose an option:", ["🔐 Login", "📝 Register"], horizontal=True, key="auth_choice")

        with st.form("login_register_form", clear_on_submit=False):
            email = st.text_input("📧 Email Address", placeholder="Enter your email address", key="email_input")
            password = st.text_input("🔒 Password", type="password", placeholder="Enter your password", key="password_input")
            st.markdown("<div style='margin: 1.5rem 0;'></div>", unsafe_allow_html=True)
            col_a, col_b = st.columns(2)
            if "🔐 Login" in choice:
                login_btn = col_a.form_submit_button("🚀 Login", use_container_width=True)
                register_btn = False
            else:
                register_btn = col_a.form_submit_button("📝 Create Account", use_container_width=True)
                login_btn = False
            st.markdown("<div style='text-align: center; color: rgba(255,255,255,0.5); margin: 1rem 0;'>── or ──</div>", unsafe_allow_html=True)
            google_signin = st.form_submit_button("🌟 Continue with Google", use_container_width=True)

    if google_signin:
        logger.info("Google sign-in initiated")
        oauth_url = get_google_oauth_url()
        if oauth_url:
            st.markdown(f'<meta http-equiv="refresh" content="1; url={oauth_url}">', unsafe_allow_html=True)
            st.link_button("🔗 Click here if redirect doesn't work automatically", oauth_url, use_container_width=True)
        else:
            st.error("❌ Failed to generate Google OAuth URL. Please check your configuration.")
        st.markdown('</div>', unsafe_allow_html=True)

    if register_btn:
        if not email or not password:
            st.error("⚠️ Please enter both email and password.")
        elif len(password) < 6:
            st.error("⚠️ Password must be at least 6 characters long.")
        elif "@" not in email or "." not in email:
            st.error("⚠️ Please enter a valid email address.")
        else:
            with st.spinner("Creating your account..."):
                success, msg = register_user(email, password)
                if success:
                    st.success(f"🎉 {msg} Please login with your credentials.")
                else:
                    st.error(f"❌ {msg}")

    if login_btn:
        if not email or not password:
            st.error("⚠️ Please enter both email and password.")
        else:
            with st.spinner("Signing you in..."):
                success, user_or_msg = login_user(email, password)
                if success:
                    st.session_state.authenticated = True
                    st.session_state.user_email = email
                    st.session_state.user = user_or_msg
                    st.session_state.messages = []
                    logger.info(f"User logged in: {email}")
                    st.success("✅ Login successful! Redirecting...")
                    st.rerun()
                else:
                    st.error(f"❌ {user_or_msg}")

# =============================================================================
# AUTHENTICATED CHAT UI
# =============================================================================
else:
    user = st.session_state.user
    logger.info(f"Rendering chat UI for user: {user['email']}")
    show_profile_in_sidebar = st.session_state.show_profile

    # -------------------------------------------------------------------------
    # SIDEBAR
    # -------------------------------------------------------------------------
    with st.sidebar:
        if show_profile_in_sidebar:
            # ---- Profile Section ----
            st.markdown("### 👤 Profile & Settings")
            if st.button("← Back to Chats", use_container_width=True):
                st.session_state.show_profile = False
                st.rerun()
            st.markdown("---")
            if user.get("picture"):
                st.image(user["picture"], width=80)
            if user.get("name"):
                st.write(f"**Name:** {user['name']}")
            st.write(f"**Email:** {user['email']}")
            if user.get("auth_method"):
                auth_icon = "🌟" if user['auth_method'] == 'google' else "📧"
                st.write(f"**Sign-in method:** {auth_icon} {user['auth_method'].title()}")
            st.markdown("---")
            tokens_used = user.get('tokens_used_today', 0)
            tokens_remaining = 2000 - tokens_used
            progress_percent = min(tokens_used / 2000, 1.0)
            st.markdown("#### 📊 Usage Statistics")
            st.metric("Tokens Used Today", f"{tokens_used:,}", f"{tokens_remaining:,} remaining")
            st.progress(progress_percent)
            st.markdown("---")
            st.markdown("#### 🔑 API Key")
            api_key_input = st.text_input("Custom API Key (optional)", value=user.get("api_key") or "", type="password", key="profile_api_key")
            col_api1, col_api2 = st.columns(2)
            if col_api1.button("💾 Save", use_container_width=True, key="save_api_profile"):
                users_col.update_one({"_id": user["_id"]}, {"$set": {"api_key": api_key_input}})
                user["api_key"] = api_key_input
                st.success("API Key saved!")
            if col_api2.button("🗑️ Delete", use_container_width=True, key="delete_api_profile"):
                users_col.update_one({"_id": user["_id"]}, {"$set": {"api_key": None}})
                user["api_key"] = None
                st.success("API Key deleted!")

        else:
            # ---- Model Selector ----
            st.markdown("### 🤖 Model")
            model_names = [f"{m['name']} — {m['description']}" for m in AVAILABLE_MODELS]
            model_ids = [m["id"] for m in AVAILABLE_MODELS]
            current_model_idx = model_ids.index(st.session_state.selected_model) if st.session_state.selected_model in model_ids else 0
            selected_label = st.selectbox(
                "Select Model",
                options=model_names,
                index=current_model_idx,
                key="model_selector",
                label_visibility="collapsed"
            )
            new_model_id = model_ids[model_names.index(selected_label)]
            if new_model_id != st.session_state.selected_model:
                st.session_state.selected_model = new_model_id
                st.rerun()

            st.markdown("---")

            # ---- Chat Mode Toggle ----
            st.markdown("### 💬 Chat Mode")
            mode_col1, mode_col2 = st.columns(2)
            if mode_col1.button(
                "🗨️ Normal" if st.session_state.chat_mode != "normal" else "✅ Normal",
                use_container_width=True,
                key="mode_normal"
            ):
                st.session_state.chat_mode = "normal"
                st.session_state.current_group_id = None
                st.rerun()
            if mode_col2.button(
                "👥 Group" if st.session_state.chat_mode != "group" else "✅ Group",
                use_container_width=True,
                key="mode_group"
            ):
                st.session_state.chat_mode = "group"
                st.rerun()

            st.markdown("---")

            # ---- Normal Chats or Group Chats ----
            if st.session_state.chat_mode == "normal":
                st.markdown("### 🗂️ Chat History")
                if st.button("➕ New Chat", use_container_width=True, key="new_chat_sidebar"):
                    new_chat_id = create_new_chat(user["_id"])
                    if new_chat_id:
                        st.session_state.current_chat_id = new_chat_id
                        st.session_state.messages = []
                        st.success("Started new chat!")
                        st.rerun()
                    else:
                        st.error("Failed to create new chat.")
                st.markdown("---")
                user_chats = get_user_chats(user["_id"])
                if user_chats:
                    for chat in user_chats:
                        chat_id = str(chat["_id"])
                        chat_title = chat.get('title', 'Untitled Chat')
                        message_count = chat.get('message_count', 0)
                        created_at = chat.get('created_at', datetime.now(timezone.utc))
                        try:
                            date_str = created_at.strftime('%Y-%m-%d %H:%M')
                        except Exception:
                            date_str = 'Unknown'
                        col_c1, col_c2 = st.columns([4, 1])
                        with col_c1:
                            if st.button(
                                f"💬 {chat_title[:25]}{'...' if len(chat_title) > 25 else ''}",
                                key=f"chat_{chat_id}",
                                help=f"Messages: {message_count}\nCreated: {date_str}",
                                use_container_width=True
                            ):
                                st.session_state.current_chat_id = chat_id
                                st.session_state.messages = get_chat_messages(chat_id)
                                st.session_state.chat_mode = "normal"
                                st.rerun()
                        with col_c2:
                            if st.button("🗑️", key=f"delete_{chat_id}", help="Delete chat"):
                                if delete_chat(chat_id):
                                    if chat_id == st.session_state.current_chat_id:
                                        st.session_state.current_chat_id = None
                                        st.session_state.messages = []
                                    st.success("Chat deleted!")
                                    st.rerun()
                else:
                    st.info("No chat history yet. Start a new conversation!")

            else:
                # ---- Group Chats ----
                st.markdown("### 👥 Group Chats")
                if st.button("➕ Create Group", use_container_width=True, key="create_group_btn"):
                    st.session_state.show_group_manager = not st.session_state.show_group_manager
                    st.rerun()

                if st.session_state.show_group_manager:
                    with st.form("create_group_form"):
                        grp_name = st.text_input("Group Name", placeholder="e.g. Study Group")
                        grp_emails = st.text_area("Member Emails (comma-separated)", placeholder="alice@example.com, bob@example.com")
                        submitted = st.form_submit_button("Create", use_container_width=True)
                        if submitted:
                            if grp_name.strip():
                                emails = [e.strip() for e in grp_emails.split(",") if e.strip()]
                                new_grp_id = create_group_chat(user["_id"], grp_name.strip(), emails)
                                if new_grp_id:
                                    st.session_state.current_group_id = new_grp_id
                                    st.session_state.messages = []
                                    st.session_state.show_group_manager = False
                                    st.success(f"Group '{grp_name}' created!")
                                    st.rerun()
                                else:
                                    st.error("Failed to create group.")
                            else:
                                st.error("Please enter a group name.")

                st.markdown("---")
                group_chats = get_user_group_chats(user["_id"])
                if group_chats:
                    for grp in group_chats:
                        grp_id = str(grp["_id"])
                        grp_title = grp.get("name", "Unnamed Group")
                        col_g1, col_g2 = st.columns([4, 1])
                        with col_g1:
                            if st.button(
                                f"👥 {grp_title[:22]}{'...' if len(grp_title) > 22 else ''}",
                                key=f"grp_{grp_id}",
                                use_container_width=True
                            ):
                                st.session_state.current_group_id = grp_id
                                st.session_state.messages = get_chat_messages(grp_id)
                                st.rerun()
                        with col_g2:
                            members_str = ", ".join(grp.get("member_emails", []))
                            st.markdown(f"<span title='{members_str}' style='color:rgba(255,255,255,0.4);font-size:0.75rem;'>👤{len(grp.get('member_emails', []))+1}</span>", unsafe_allow_html=True)
                else:
                    st.info("No group chats yet.")

            st.markdown("---")

            # ---- Document Upload ----
            st.markdown("### 📎 Documents")
            support_info = []
            if PDF_SUPPORT:
                support_info.append("pdf")
            if DOCX_SUPPORT:
                support_info.append("docx")
            if EXCEL_SUPPORT:
                support_info.append("xlsx, xls")
            accepted_types = ["pdf", "docx", "xlsx", "xls"] if support_info else []

            if accepted_types:
                uploaded_file = st.file_uploader(
                    "Upload document",
                    type=accepted_types,
                    key="doc_uploader",
                    label_visibility="collapsed"
                )
                if uploaded_file is not None:
                    already = any(d["filename"] == uploaded_file.name for d in st.session_state.uploaded_docs)
                    if not already:
                        with st.spinner(f"Processing {uploaded_file.name}..."):
                            doc_data = extract_text_from_file(uploaded_file)
                            st.session_state.uploaded_docs.append(doc_data)
                        st.success(f"✅ Added: {uploaded_file.name}")
                        st.rerun()

                if st.session_state.uploaded_docs:
                    st.markdown("**Loaded documents:**")
                    for idx, doc in enumerate(st.session_state.uploaded_docs):
                        col_d1, col_d2 = st.columns([4, 1])
                        with col_d1:
                            st.markdown(f'<span class="doc-chip">📄 {doc["filename"][:20]}</span>', unsafe_allow_html=True)
                        with col_d2:
                            if st.button("✕", key=f"remove_doc_{idx}", help="Remove document"):
                                st.session_state.uploaded_docs.pop(idx)
                                st.rerun()
            else:
                st.info("Install pypdf, python-docx, openpyxl for document support.")

            st.markdown("---")

            # ---- Web Search Toggle ----
            st.markdown("### 🔍 Web Search")
            if WEB_SEARCH_SUPPORT:
                ws_enabled = st.checkbox(
                    "Enable web search",
                    value=st.session_state.web_search_enabled,
                    key="web_search_checkbox"
                )
                if ws_enabled != st.session_state.web_search_enabled:
                    st.session_state.web_search_enabled = ws_enabled
                    st.rerun()
                if st.session_state.web_search_enabled:
                    st.markdown('<span style="color:#68d391;">● Web search active</span>', unsafe_allow_html=True)
            else:
                st.info("Install duckduckgo_search for web search support.")

            st.markdown("---")

            # ---- Quick Stats ----
            if st.session_state.chat_mode == "normal":
                user_chats_count = len(get_user_chats(user["_id"]))
                st.markdown(f"**Total Chats:** {user_chats_count}")
                if st.session_state.messages:
                    user_msg_count = len([m for m in st.session_state.messages if m[0] == 'user'])
                    st.markdown(f"**Current Chat:** {user_msg_count} messages")

    # -------------------------------------------------------------------------
    # MAIN CONTENT
    # -------------------------------------------------------------------------

    # Determine layout: canvas or full-width
    if st.session_state.show_canvas and st.session_state.canvas_content:
        chat_col, canvas_col = st.columns([6, 4])
    else:
        chat_col = st.container()
        canvas_col = None

    # ---- Canvas Panel ----
    if canvas_col is not None:
        with canvas_col:
            st.markdown('<div class="canvas-panel">', unsafe_allow_html=True)
            st.markdown("### 🎨 Canvas")
            canvas_lang = st.selectbox(
                "Language",
                CANVAS_LANGUAGES,
                index=CANVAS_LANGUAGES.index(st.session_state.canvas_language) if st.session_state.canvas_language in CANVAS_LANGUAGES else 0,
                key="canvas_lang_selector"
            )
            st.session_state.canvas_language = canvas_lang
            st.code(st.session_state.canvas_content, language=canvas_lang)
            if st.button("✕ Close Canvas", use_container_width=True, key="close_canvas"):
                st.session_state.show_canvas = False
                st.session_state.canvas_content = None
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)

    # ---- Chat Area ----
    def render_chat_area():
        inner_col1, inner_col2, inner_col3 = st.columns([1, 6, 1])
        with inner_col2:
            # Group chat header
            if st.session_state.chat_mode == "group" and st.session_state.current_group_id:
                try:
                    grp_doc = db.group_chats.find_one({"_id": ObjectId(st.session_state.current_group_id)})
                    if grp_doc:
                        members_display = ", ".join(grp_doc.get("member_emails", []))
                        st.markdown(f"""
                            <div style="background:rgba(102,126,234,0.1);border:1px solid rgba(102,126,234,0.3);
                            border-radius:12px;padding:1rem;margin-bottom:1rem;">
                                <strong style="color:#667eea;">👥 {grp_doc.get('name','Group Chat')}</strong><br>
                                <span style="color:rgba(255,255,255,0.6);font-size:0.85rem;">Members: {members_display or 'Just you'}</span>
                            </div>
                        """, unsafe_allow_html=True)
                        # Add member form
                        with st.expander("➕ Add Member"):
                            new_member_email = st.text_input("Member email", key="add_member_email")
                            if st.button("Add", key="add_member_btn"):
                                if new_member_email:
                                    if add_member_to_group(st.session_state.current_group_id, new_member_email):
                                        st.success(f"Added {new_member_email}")
                                        st.rerun()
                                    else:
                                        st.error("Failed to add member.")
                except Exception as e:
                    logger.error(f"Error loading group: {e}")

            # Welcome screen for empty chat
            if not st.session_state.messages:
                st.markdown('<div style="text-align: center; margin: 3rem 0 4rem 0; opacity: 0.9;">', unsafe_allow_html=True)
                st.markdown('<div style="font-size: 3rem; margin-bottom: 1rem;">🤖</div>', unsafe_allow_html=True)
                st.markdown('<h1 style="font-size: 2.2rem; margin-bottom: 1rem; font-weight: 600; color: #fff;">YAHANAR AI</h1>', unsafe_allow_html=True)
                st.markdown('<p style="color: rgba(255,255,255,0.7); font-size: 1.1rem; margin-bottom: 2rem;">How can I help you today?</p>', unsafe_allow_html=True)
                fc1, fc2, fc3, fc4 = st.columns(4)
                cards = [
                    ("💡", "Creative Writing", "Stories, essays, and creative content"),
                    ("💻", "Programming", "Debug code and learn programming"),
                    ("🔍", "Research", "Deep dive into topics and analysis"),
                    ("🎨", "Creative Ideas", "Brainstorm innovative solutions"),
                ]
                for col, (icon, title, desc) in zip([fc1, fc2, fc3, fc4], cards):
                    with col:
                        st.markdown(f"""
                            <div style="background: rgba(255,255,255,0.05); border-radius: 12px; padding: 20px;
                            text-align: left; border: 1px solid rgba(255,255,255,0.1); height: 120px;">
                                <div style="font-size: 1.5rem; margin-bottom: 8px;">{icon}</div>
                                <h3 style="margin: 0 0 8px 0; font-size: 1rem; color: #fff;">{title}</h3>
                                <p style="margin: 0; color: rgba(255,255,255,0.6); font-size: 0.85rem;">{desc}</p>
                            </div>
                        """, unsafe_allow_html=True)
                st.markdown('</div>', unsafe_allow_html=True)

            # Display messages
            logger.info("Rendering chat history")
            for i, (role, text) in enumerate(st.session_state.messages):
                if role == "user":
                    st.markdown(f"""
                        <div class="message-container user-container fade-in">
                            <div class="user-message">{text}</div>
                            <div class="message-avatar user-avatar">👤</div>
                        </div>
                    """, unsafe_allow_html=True)
                else:
                    st.markdown(f'''
                        <div class="message-container assistant-container slide-in">
                            <div class="message-avatar assistant-avatar">🤖</div>
                            <div class="assistant-message"><div class="message-content">{text}</div></div>
                        </div>
                    ''', unsafe_allow_html=True)

            if st.session_state.messages:
                st.markdown('<div style="margin-bottom: 2rem;"></div>', unsafe_allow_html=True)

            # Feature toolbar
            current_model_name = next((m["name"] for m in AVAILABLE_MODELS if m["id"] == st.session_state.selected_model), st.session_state.selected_model)
            ws_indicator = "🟢" if st.session_state.web_search_enabled else "⚪"
            docs_indicator = f"📎{len(st.session_state.uploaded_docs)}" if st.session_state.uploaded_docs else "📎"
            canvas_indicator = "🎨✓" if st.session_state.show_canvas else "🎨"

            toolbar_col1, toolbar_col2, toolbar_col3, toolbar_col4 = st.columns([1, 1, 1, 3])
            with toolbar_col1:
                st.markdown(f'<div style="padding:6px;text-align:center;color:rgba(255,255,255,0.7);">{docs_indicator}</div>', unsafe_allow_html=True)
            with toolbar_col2:
                st.markdown(f'<div style="padding:6px;text-align:center;" title="Web Search">{ws_indicator} Search</div>', unsafe_allow_html=True)
            with toolbar_col3:
                if st.button(canvas_indicator, key="toggle_canvas_toolbar", help="Toggle Canvas"):
                    st.session_state.show_canvas = not st.session_state.show_canvas
                    st.rerun()
            with toolbar_col4:
                st.markdown(f'<div style="padding:6px;color:rgba(255,255,255,0.5);font-size:0.85rem;">Model: <span class="model-badge">{current_model_name}</span></div>', unsafe_allow_html=True)

    if canvas_col is not None:
        with chat_col:
            render_chat_area()
    else:
        render_chat_area()

    # ---- Chat Input ----
    is_group = st.session_state.chat_mode == "group"
    placeholder_text = "Message group..." if is_group else "Message Krutrim AI..."
    if is_group and not st.session_state.current_group_id:
        st.info("👥 Select or create a group chat from the sidebar to start chatting.")
    else:
        prompt = st.chat_input(placeholder_text, key="chat_input")

        if prompt:
            logger.info(f"User sent message: {prompt}")
            tokens_needed = 512

            if not can_use_tokens(user, tokens_needed) and not user.get("api_key"):
                logger.warning("Token limit reached for user")
                st.error("""
                    🚫 **Daily Token Limit Reached**

                    You've used all 2,000 free tokens for today. To continue:
                    • 🔑 Add your Krutrim API key in the sidebar
                    • ⏰ Wait for tomorrow's reset
                """, icon="⚠️")
            else:
                st.session_state.messages.append(("user", prompt))

                with st.status("✨ Generating response...", expanded=False) as status:
                    try:
                        # Build system context
                        system_parts = []
                        if st.session_state.uploaded_docs:
                            doc_ctx = get_document_context(st.session_state.uploaded_docs)
                            if doc_ctx:
                                system_parts.append(f"The user has uploaded the following documents for reference:\n\n{doc_ctx}")

                        if st.session_state.web_search_enabled:
                            with st.spinner("🔍 Searching the web..."):
                                search_results = web_search(prompt)
                            if search_results:
                                system_parts.append(f"Web search results for the user's query:\n\n{search_results}")

                        messages_payload = [{"role": r, "content": c} for r, c in st.session_state.messages]

                        if system_parts:
                            system_content = "\n\n".join(system_parts)
                            messages_payload = [{"role": "system", "content": system_content}] + messages_payload

                        reply, used_tokens = chat_with_krutrim(
                            messages_payload,
                            api_key=user.get("api_key"),
                            model=st.session_state.selected_model,
                            max_tokens=2048
                        )

                        st.session_state.messages.append(("assistant", reply))

                        # Extract code blocks → canvas
                        code_blocks = extract_code_blocks(reply)
                        if code_blocks:
                            lang, code = code_blocks[0]
                            st.session_state.canvas_content = code
                            st.session_state.canvas_language = lang if lang != "text" else "python"
                            st.session_state.show_canvas = True

                        # Save to database
                        logger.info("Saving chat messages to database")
                        try:
                            chat_id_to_use = st.session_state.current_group_id if is_group else st.session_state.current_chat_id
                            if not chat_id_to_use:
                                if is_group:
                                    st.error("No group chat selected.")
                                    st.stop()
                                else:
                                    new_chat_id = create_new_chat(user["_id"])
                                    if new_chat_id:
                                        st.session_state.current_chat_id = new_chat_id
                                        chat_id_to_use = new_chat_id
                                    else:
                                        st.error("Failed to create chat.")
                                        st.stop()
                            save_chat_message(user["_id"], chat_id_to_use, "user", prompt)
                            save_chat_message(user["_id"], chat_id_to_use, "assistant", reply)
                        except Exception as e:
                            logger.error(f"Failed to save chat to database: {e}")

                        # Update token usage
                        increment_tokens(user, used_tokens)
                        user = users_col.find_one({"_id": user["_id"]})
                        st.session_state.user = user
                        CHAT_COUNT.inc()
                        status.update(label="✅ Response generated!", state="complete")

                    except Exception as e:
                        logger.error(f"Error generating response: {e}")
                        st.session_state.messages.append(("assistant", "❌ Sorry, I encountered an error. Please try again."))
                        status.update(label="❌ Error occurred", state="error")

                st.rerun()

    # Auto-scroll
    if st.session_state.messages:
        st.markdown("""
            <script>
                setTimeout(function() {
                    window.scrollTo({ top: document.body.scrollHeight, behavior: 'smooth' });
                }, 100);
            </script>
        """, unsafe_allow_html=True)

    # Footer
    st.markdown("---")
    st.markdown("""
        <div style="text-align: center; color: rgba(255,255,255,0.5); padding: 2rem 0;">
            <p>🤖 Powered by <strong>Krutrim AI</strong> | 🔒 Your privacy is protected</p>
            <p style="font-size: 0.8rem;">Built with ❤️ using Streamlit | 📊 <a href="http://localhost:8000" style="color: #667eea;">Metrics Dashboard</a></p>
        </div>
    """, unsafe_allow_html=True)
